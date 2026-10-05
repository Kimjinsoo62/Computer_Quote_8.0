# -*- coding: utf-8 -*-
"""컴퓨터 작업내역서 관리 및 엑셀 내보내기 모듈"""
from __future__ import annotations

import json
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.quote_info import load_quote_info

DATA_FILE = Path(__file__).resolve().parent.parent / "task_reports.json"
TEMPLATE_FILE = Path(__file__).resolve().parent.parent / "컴퓨터_작업내역서.xlsx"


def load_task_reports() -> list[dict[str, Any]]:
    """저장된 작업내역서 목록을 반환한다."""
    if DATA_FILE.exists():
        try:
            return json.loads(DATA_FILE.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            pass
    return []


def save_all_task_reports(reports: list[dict[str, Any]]) -> None:
    """작업내역서 전체 목록을 파일에 저장한다."""
    DATA_FILE.write_text(
        json.dumps(reports, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def get_task_report(report_id: str) -> dict[str, Any] | None:
    """특정 ID의 작업내역서를 조회한다."""
    reports = load_task_reports()
    for rep in reports:
        if rep.get("id") == report_id:
            return rep
    return None


def create_task_report(data: dict[str, Any]) -> dict[str, Any]:
    """새 작업내역서를 등록한다."""
    reports = load_task_reports()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    report_id = str(uuid.uuid4())[:8]
    qi = load_quote_info()

    report = {
        "id": report_id,
        "created_at": now_str,
        "updated_at": now_str,
        "company_name": data.get("company_name") or qi.get("supplier_name", "대한민국주식회사"),
        "company_address": data.get("company_address") or qi.get("supplier_address", "서울시"),
        "company_phone": data.get("company_phone") or qi.get("supplier_phone", "010-1234-5678"),
        "business_number": data.get("business_number", "121-18-24250"),
        "doc_no": data.get("doc_no") or f"TR-{datetime.now().strftime('%Y%m%d')}-{report_id[:4].upper()}",
        "receipt_date": data.get("receipt_date") or datetime.now().strftime("%Y-%m-%d"),
        
        # 1. 고객 정보
        "customer_name": data.get("customer_name", ""),
        "customer_phone": data.get("customer_phone", ""),
        "customer_address": data.get("customer_address", ""),
        "customer_email": data.get("customer_email", ""),
        "contact_method": data.get("contact_method", "전화"),  # 전화, 문자, 이메일
        
        # 2. 장비 정보
        "device_type": data.get("device_type", "데스크탑"),  # 데스크탑, 노트북, 서버, 기타
        "device_model": data.get("device_model", ""),
        "serial_number": data.get("serial_number", ""),
        "purchase_date": data.get("purchase_date", ""),
        "accessories": data.get("accessories", []),  # ['충전기', '마우스', '케이블', '가방']
        "accessories_custom": data.get("accessories_custom", ""),
        
        # 3. 증상 / 요청 사항
        "symptoms": data.get("symptoms", ""),
        
        # 4. 진단 및 작업 내역
        "tasks": data.get("tasks", [
            {"no": 1, "description": "", "parts": "", "cost": 0},
            {"no": 2, "description": "", "parts": "", "cost": 0},
            {"no": 3, "description": "", "parts": "", "cost": 0},
        ]),
        "total_cost": int(data.get("total_cost", 0)),
        
        # 5. 처리 현황
        "status": data.get("status", "완료"),  # 접수, 진단중, 수리중, 완료, 인도완료
        "expected_date": data.get("expected_date", datetime.now().strftime("%Y-%m-%d")),
        "technician": data.get("technician", "김대건"),
        "warranty_period": data.get("warranty_period", "3개월"),
        
        # 6. 확인 및 서명
        "customer_sign": data.get("customer_sign", ""),
        "technician_sign": data.get("technician_sign", ""),
    }

    # 계산된 총액
    total = sum(int(t.get("cost", 0) or 0) for t in report["tasks"])
    if total > 0 or not report["total_cost"]:
        report["total_cost"] = total

    reports.insert(0, report)
    save_all_task_reports(reports)
    return report


def update_task_report(report_id: str, data: dict[str, Any]) -> dict[str, Any] | None:
    """작업내역서를 수정한다."""
    reports = load_task_reports()
    target_idx = -1
    for idx, rep in enumerate(reports):
        if rep.get("id") == report_id:
            target_idx = idx
            break

    if target_idx == -1:
        return None

    existing = reports[target_idx]
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    existing.update(data)
    existing["id"] = report_id
    existing["updated_at"] = now_str

    # 총액 갱신
    if "tasks" in existing:
        existing["total_cost"] = sum(int(t.get("cost", 0) or 0) for t in existing["tasks"])

    reports[target_idx] = existing
    save_all_task_reports(reports)
    return existing


def delete_task_report(report_id: str) -> bool:
    """작업내역서를 삭제한다."""
    reports = load_task_reports()
    initial_len = len(reports)
    reports = [rep for rep in reports if rep.get("id") != report_id]
    if len(reports) < initial_len:
        save_all_task_reports(reports)
        return True
    return False


def build_task_report_excel(report: dict[str, Any]) -> BytesIO:
    """작업내역서 엑셀 파일을 생성하여 BytesIO로 반환한다."""
    if TEMPLATE_FILE.exists():
        wb = openpyxl.load_workbook(TEMPLATE_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "작업내역서"

    # 1. 헤더 영역
    comp_name = report.get("company_name", "0000 ")
    ws["A1"] = f"{comp_name} 컴퓨터 판매·수리센터"
    ws["A3"] = f"주소 : {report.get('company_address', '')}"
    ws["A4"] = f"전화 : {report.get('company_phone', '')}  |   사업자등록번호: {report.get('business_number', '')}"
    
    ws["E3"] = f"접수번호 : {report.get('doc_no', '')}"
    r_date = report.get("receipt_date", "")
    ws["E4"] = f"접수일자 : {r_date}"

    # 2. 고객 정보
    ws["B7"] = report.get("customer_name", "")
    ws["E7"] = report.get("customer_phone", "")
    ws["B8"] = report.get("customer_address", "")
    ws["B9"] = report.get("customer_email", "")

    contact_m = report.get("contact_method", "전화")
    ws["E9"] = f"{'■' if contact_m == '전화' else '□'} 전화   {'■' if contact_m == '문자' else '□'} 문자   {'■' if contact_m == '이메일' else '□'} 이메일"

    # 3. 장비 정보
    dev_t = report.get("device_type", "데스크탑")
    ws["B11"] = f"{'■' if dev_t == '데스크탑' else '□'} 데스크탑   {'■' if dev_t == '노트북' else '□'} 노트북   {'■' if dev_t == '서버' else '□'} 서버   {'■' if dev_t == '기타' else '□'} 기타"
    ws["E11"] = report.get("device_model", "")
    ws["B12"] = report.get("serial_number", "")
    ws["E12"] = report.get("purchase_date", "")

    acc_list = report.get("accessories", [])
    acc_str = (
        f"{'■' if '충전기' in acc_list else '□'} 충전기   "
        f"{'■' if '마우스' in acc_list else '□'} 마우스   "
        f"{'■' if '케이블' in acc_list else '□'} 케이블   "
        f"{'■' if '가방' in acc_list else '□'} 가방   "
        f"{'■' if report.get('accessories_custom') else '□'} 기타 : {report.get('accessories_custom', '')}"
    )
    ws["B13"] = acc_str

    # 4. 증상 / 요청 사항
    ws["A15"] = report.get("symptoms", "")

    # 5. 진단 및 작업 내역 (Row 20, 21, 22 등)
    tasks = report.get("tasks", [])
    for idx in range(3):
        row_num = 20 + idx
        if idx < len(tasks):
            t = tasks[idx]
            ws[f"A{row_num}"] = idx + 1
            ws[f"B{row_num}"] = t.get("description", "")
            ws[f"D{row_num}"] = t.get("parts", "")
            cost_val = int(t.get("cost", 0) or 0)
            ws[f"F{row_num}"] = cost_val if cost_val > 0 else ""
            if cost_val > 0:
                ws[f"F{row_num}"].number_format = "#,##0"
        else:
            ws[f"A{row_num}"] = idx + 1
            ws[f"B{row_num}"] = ""
            ws[f"D{row_num}"] = ""
            ws[f"F{row_num}"] = ""

    ws["F23"] = int(report.get("total_cost", 0) or 0)
    ws["F23"].number_format = "#,##0"

    # 6. 처리 현황
    status_val = report.get("status", "완료")
    status_str = (
        f"{'■' if status_val == '접수' else '□'} 접수   "
        f"{'■' if status_val == '진단중' else '□'} 진단중   "
        f"{'■' if status_val == '수리중' else '□'} 수리중   "
        f"{'■' if status_val == '완료' else '□'} 완료   "
        f"{'■' if status_val == '인도완료' else '□'} 인도완료"
    )
    ws["B25"] = status_str
    ws["E25"] = report.get("expected_date", "")
    ws["B26"] = report.get("technician", "김대건")
    ws["E26"] = report.get("warranty_period", "3개월")

    # 7. 확인 및 서명
    if report.get("customer_sign"):
        ws["A29"] = f"서명: {report.get('customer_sign')}"
    if report.get("technician_sign"):
        ws["D29"] = f"서명: {report.get('technician_sign')}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
